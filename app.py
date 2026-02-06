from __future__ import annotations

from dataclasses import dataclass, asdict
import csv
import importlib.util
import io
import json
import logging
import os
import re
import secrets
import socket
import subprocess
import sys
import threading
import webbrowser
from typing import Any
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from docx import Document
import requests
from flask import Flask, Response, render_template, request, session
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", secrets.token_hex(16))

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

DEFAULT_TIMEOUT = 20
REQUEST_RETRIES = 2
MAX_RESULTS = 80
MAX_TEXT_LENGTH = 4000
SCRIPT_SCAN_LIMIT = 30000
MAX_DATA_SOURCES = 60

AUTO_INSTALL = os.getenv("AUTO_INSTALL_DEPS", "1") == "1"
REQUIRED_MODULES = ("flask", "requests", "bs4", "openpyxl", "playwright", "docx")


def _missing_modules() -> list[str]:
    missing = []
    for module_name in REQUIRED_MODULES:
        if importlib.util.find_spec(module_name) is None:
            missing.append(module_name)
    return missing


def _auto_install_requirements() -> None:
    if not AUTO_INSTALL:
        return
    missing = _missing_modules()
    if not missing:
        return
    logger.info("Detected missing modules: %s. Installing requirements...", ", ".join(missing))
    subprocess.run(
        [sys.executable, "-m", "pip", "install", "-r", "requirements.txt"],
        check=False,
    )

CHART_KEYWORDS = (
    "echarts",
    "highcharts",
    "chart.js",
    "chartjs",
    "d3",
    "plotly",
    "apexcharts",
    "recharts",
    "series",
    "dataset",
    "datasets",
    "xAxis",
    "yAxis",
    "values",
    "labels",
    "option",
    "options",
    "config",
)

DATA_PATTERN = re.compile(
    r"(?P<key>series|dataset|datasets|xAxis|yAxis|values|labels|data|option|options|config)\s*:\s*(?P<value>[\[{])",
    re.IGNORECASE,
)

SCRIPT_JSON_TYPES = {"application/json", "application/ld+json"}


@dataclass
class DataSource:
    label: str
    value: str
    formatted: str | None
    source: str


def _strip_js_comments(text: str) -> str:
    text = re.sub(r"/\*.*?\*/", "", text, flags=re.DOTALL)
    text = re.sub(r"//.*?$", "", text, flags=re.MULTILINE)
    return text


def _find_balanced_block(text: str, start_index: int) -> tuple[str, int] | None:
    stack = []
    in_string = False
    string_char = ""
    escape = False

    for index in range(start_index, len(text)):
        char = text[index]
        if escape:
            escape = False
            continue
        if char == "\\":
            escape = True
            continue
        if in_string:
            if char == string_char:
                in_string = False
            continue
        if char in ("'", '"'):
            in_string = True
            string_char = char
            continue
        if char in "[{":
            stack.append(char)
        elif char in "]}":
            if not stack:
                break
            opener = stack.pop()
            if (opener == "[" and char != "]") or (opener == "{" and char != "}"):
                break
            if not stack:
                return text[start_index : index + 1], index + 1
    return None


def _normalize_to_json(raw: str) -> str | None:
    cleaned = raw.strip()
    try:
        parsed = json.loads(cleaned)
        return json.dumps(parsed, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        pass

    normalized = cleaned.replace("'", '"')
    normalized = normalized.replace("undefined", "null").replace("NaN", "null")
    normalized = normalized.replace("true", "true").replace("false", "false")
    try:
        parsed = json.loads(normalized)
        return json.dumps(parsed, ensure_ascii=False, indent=2)
    except json.JSONDecodeError:
        return None


def normalize_url(raw_url: str) -> str:
    parsed = urlparse(raw_url)
    if not parsed.scheme:
        return f"https://{raw_url}"
    return raw_url


def _is_private_host(hostname: str) -> bool:
    if hostname in {"localhost", "127.0.0.1", "::1"}:
        return True
    private_prefixes = ("10.", "192.168.", "172.16.", "172.17.", "172.18.", "172.19.")
    if hostname.startswith(private_prefixes):
        return True
    try:
        ip = socket.gethostbyname(hostname)
    except socket.gaierror:
        return False
    return ip.startswith(private_prefixes) or ip.startswith("127.")


def _parse_cookie_string(cookie_str: str) -> dict[str, str]:
    cookies: dict[str, str] = {}
    for pair in cookie_str.split(";"):
        if "=" in pair:
            key, value = pair.split("=", 1)
            cookies[key.strip()] = value.strip()
    return cookies


def _parse_headers(headers_str: str) -> dict[str, str] | None:
    if not headers_str:
        return {}
    try:
        parsed = json.loads(headers_str)
    except json.JSONDecodeError:
        return None
    return {str(key): str(value) for key, value in parsed.items()}

def _build_session() -> requests.Session:
    retry = Retry(
        total=REQUEST_RETRIES,
        backoff_factor=0.5,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=("GET",),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry)
    session = requests.Session()
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    return session


def _extract_from_scripts(
    soup: BeautifulSoup, custom_regex: str | None
) -> tuple[list[DataSource], list[str]]:
    data_sources: list[DataSource] = []
    raw_snippets: list[str] = []

    for script in soup.find_all("script"):
        script_type = (script.get("type") or "").lower()
        content = script.string or script.get_text() or ""
        if not content:
            continue

        if script_type in SCRIPT_JSON_TYPES:
            try:
                parsed = json.loads(content)
                formatted = json.dumps(parsed, ensure_ascii=False, indent=2)
                data_sources.append(
                    DataSource(
                        label="脚本 JSON 数据",
                        value=formatted[:MAX_TEXT_LENGTH],
                        formatted=formatted,
                        source="script[type=application/json]",
                    )
                )
            except json.JSONDecodeError:
                pass

        lowered = content.lower()
        if any(keyword in lowered for keyword in CHART_KEYWORDS):
            snippet = " ".join(content.split())
            raw_snippets.append(snippet[:SCRIPT_SCAN_LIMIT])

            cleaned = _strip_js_comments(content)
            for match in DATA_PATTERN.finditer(cleaned):
                key = match.group("key")
                block = _find_balanced_block(cleaned, match.start("value"))
                if not block:
                    continue
                value = block[0].strip()
                formatted = _normalize_to_json(value)
                label = f"{key} 数据"
                data_sources.append(
                    DataSource(
                        label=label,
                        value=value[:MAX_TEXT_LENGTH],
                        formatted=formatted,
                        source="script",
                    )
                )
                if len(data_sources) >= MAX_DATA_SOURCES:
                    break

        if custom_regex:
            try:
                regex = re.compile(custom_regex, re.IGNORECASE | re.DOTALL)
                for idx, match in enumerate(regex.finditer(content), start=1):
                    value = match.group(0)
                    formatted = _normalize_to_json(value)
                    label = f"自定义匹配 #{idx}"
                    data_sources.append(
                        DataSource(
                            label=label,
                            value=value[:MAX_TEXT_LENGTH],
                            formatted=formatted,
                            source="custom regex",
                        )
                    )
            except re.error:
                pass

        if len(data_sources) >= MAX_DATA_SOURCES:
            break

    return data_sources[:MAX_DATA_SOURCES], raw_snippets[:MAX_RESULTS]


def _extract_tables(soup: BeautifulSoup) -> list[DataSource]:
    tables = soup.find_all("table")
    if not tables:
        return []

    best_table: list[list[str]] = []
    best_row_count = 0

    for table in tables:
        thead = table.find("thead")
        tbody = table.find("tbody")
        header_rows = thead.find_all("tr") if thead else []
        body_rows = tbody.find_all("tr") if tbody else []

        if not header_rows:
            header_rows = table.find_all("tr")[:1]
        if not body_rows:
            body_rows = table.find_all("tr")[1:]

        if not header_rows:
            continue

        header_cells = header_rows[0].find_all(["th", "td"])
        headers = [cell.get_text(strip=True) or f"列{idx+1}" for idx, cell in enumerate(header_cells)]
        if not headers:
            continue

        table_matrix: list[list[str]] = [headers]
        for row in body_rows:
            cells = [cell.get_text(strip=True) for cell in row.find_all(["td", "th"])]
            if not any(cells):
                continue
            padded = cells + [""] * (len(headers) - len(cells))
            table_matrix.append(padded[: len(headers)])

        row_count = len(table_matrix) - 1
        if row_count > best_row_count:
            best_row_count = row_count
            best_table = table_matrix

    if best_table:
        formatted = json.dumps(best_table, ensure_ascii=False, indent=2)
        return [
            DataSource(
                label="主表格数据",
                value=formatted[:MAX_TEXT_LENGTH],
                formatted=formatted,
                source="table",
            )
        ]
    return []


def _extract_svg_data(soup: BeautifulSoup) -> list[DataSource]:
    results: list[DataSource] = []
    svg_elements = soup.find_all("svg")
    for index, svg in enumerate(svg_elements, start=1):
        entries = []
        for node in svg.find_all(["path", "rect", "circle", "text"]):
            entry = {"tag": node.name}
            for attr in ("d", "x", "y", "cx", "cy", "r", "width", "height", "transform"):
                if node.has_attr(attr):
                    entry[attr] = node.get(attr)
            if node.name == "text":
                entry["text"] = node.get_text(strip=True)
            if entry:
                entries.append(entry)
        if entries:
            formatted = json.dumps(entries, ensure_ascii=False, indent=2)
            results.append(
                DataSource(
                    label=f"SVG 图表元素 #{index}",
                    value=formatted[:MAX_TEXT_LENGTH],
                    formatted=formatted,
                    source="svg",
                )
            )
    return results


def extract_chart_data(
    html: str, custom_regex: str | None
) -> tuple[list[DataSource], list[str]]:
    soup = BeautifulSoup(html, "html.parser")
    data_sources, raw_snippets = _extract_from_scripts(soup, custom_regex)
    data_sources.extend(_extract_tables(soup))
    data_sources.extend(_extract_svg_data(soup))
    return data_sources[:MAX_DATA_SOURCES], raw_snippets


def _build_requests_options(form: dict[str, str]) -> tuple[dict[str, str], dict[str, str], dict[str, str], int, str | None]:
    headers = {"User-Agent": "Mozilla/5.0"}
    cookies: dict[str, str] = {}
    proxies: dict[str, str] = {}
    proxy = form.get("proxy", "").strip()
    timeout_raw = form.get("timeout", "").strip()
    custom_regex = form.get("custom_regex", "").strip() or None

    headers_input = form.get("headers", "").strip()
    parsed_headers = _parse_headers(headers_input)
    if parsed_headers is None:
        raise ValueError("Headers 必须是 JSON 格式，例如 {\"Authorization\": \"Bearer ...\"}。")
    headers.update(parsed_headers)

    cookie_input = form.get("cookies", "").strip()
    if cookie_input:
        cookies.update(_parse_cookie_string(cookie_input))

    if proxy:
        proxies = {"http": proxy, "https": proxy}

    timeout = DEFAULT_TIMEOUT
    if timeout_raw:
        try:
            timeout = max(5, min(60, int(timeout_raw)))
        except ValueError:
            raise ValueError("超时时间必须是数字（5-60）。")

    return headers, cookies, proxies, timeout, custom_regex


def scrape_content(
    url: str,
    headers: dict[str, str],
    cookies: dict[str, str],
    proxies: dict[str, str],
    timeout: int,
    custom_regex: str | None,
) -> tuple[list[DataSource], list[str]]:
    session_client = _build_session()
    response = session_client.get(
        url,
        headers=headers,
        cookies=cookies,
        proxies=proxies,
        timeout=timeout,
    )
    response.raise_for_status()
    return extract_chart_data(response.text, custom_regex)


def _scrape_dynamic_content(
    url: str,
    headers: dict[str, str],
    cookies: dict[str, str],
    timeout: int,
    custom_regex: str | None,
    interactive: bool,
) -> tuple[list[DataSource], list[str]]:
    try:
        from playwright.sync_api import sync_playwright  # type: ignore
    except ImportError as exc:
        raise RuntimeError("未检测到 Playwright，请先安装并执行 playwright install。") from exc

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not interactive)
        context = browser.new_context()
        if headers:
            context.set_extra_http_headers(headers)
        if cookies:
            cookie_list = [
                {"name": key, "value": value, "domain": urlparse(url).hostname or ""}
                for key, value in cookies.items()
            ]
            context.add_cookies(cookie_list)
        page = context.new_page()
        api_sources: list[DataSource] = []

        def _contains_large_numeric_arrays(payload: object) -> bool:
            try:
                if isinstance(payload, list):
                    return any(
                        isinstance(item, list)
                        and len(item) >= 5
                        and all(isinstance(v, (int, float)) for v in item if v is not None)
                        for item in payload
                    )
                if isinstance(payload, dict):
                    for value in payload.values():
                        if isinstance(value, list) and len(value) >= 5:
                            if all(isinstance(v, (int, float)) for v in value if v is not None):
                                return True
                            if any(
                                isinstance(item, list)
                                and len(item) >= 5
                                and all(isinstance(v, (int, float)) for v in item if v is not None)
                                for item in value
                            ):
                                return True
                return False
            except Exception:
                return False

        def handle_response(response) -> None:
            try:
                content_type = response.headers.get("content-type", "")
                if "application/json" not in content_type:
                    return
                payload = response.json()
                if not _contains_large_numeric_arrays(payload):
                    return
                formatted = json.dumps(payload, ensure_ascii=False, indent=2)
                api_sources.append(
                    DataSource(
                        label=f"API 数组数据: {response.url}",
                        value=formatted[:MAX_TEXT_LENGTH],
                        formatted=formatted,
                        source=f"api_json:{response.url}",
                    )
                )
            except Exception:
                return

        page.on("response", handle_response)
        page.goto(url, wait_until="domcontentloaded", timeout=timeout * 1000)
        page.wait_for_load_state("networkidle", timeout=timeout * 1000)
        if interactive:
            page.wait_for_timeout(30000)
            page.wait_for_load_state("networkidle", timeout=timeout * 1000)

        chart_configs = page.evaluate(
            """
            () => {
              const results = [];
              if (typeof echarts !== "undefined") {
                const instances = [];
                document.querySelectorAll("*").forEach((el) => {
                  try {
                    const instance = echarts.getInstanceByDom(el);
                    if (instance) {
                      instances.push(instance);
                    }
                  } catch (e) {}
                });
                instances.forEach((instance, index) => {
                  try {
                    const option = instance.getOption ? instance.getOption() : null;
                    if (option) {
                      results.push({
                        type: "echarts_option",
                        label: `ECharts 配置 #${index + 1}`,
                        data: option
                      });
                      const xAxis = option.xAxis ? option.xAxis[0]?.data || option.xAxis.data : null;
                      const series = option.series || [];
                      const seriesData = series.map((item) => item?.data).filter(Boolean);
                      results.push({
                        type: "echarts_data",
                        label: `ECharts 数据 #${index + 1}`,
                        data: { xAxis, series: seriesData }
                      });
                    }
                  } catch (e) {}
                });
              }
              return results;
            }
            """
        )

        html = page.content()
        browser.close()

    data_sources, raw_snippets = extract_chart_data(html, custom_regex)
    existing_values = {source.value for source in data_sources}
    for config in chart_configs or []:
        try:
            formatted = json.dumps(config["data"], ensure_ascii=False, indent=2)
        except (TypeError, json.JSONDecodeError, KeyError):
            continue
        if formatted in existing_values:
            continue
        data_sources.append(
            DataSource(
                label=config.get("label", "图表实例"),
                value=formatted[:MAX_TEXT_LENGTH],
                formatted=formatted,
                source=f"javascript_{config.get('type', 'runtime')}",
            )
        )
    if api_sources:
        data_sources = api_sources + data_sources
    return data_sources, raw_snippets


def _serialize_data_sources(data_sources: list[DataSource]) -> str:
    return json.dumps([asdict(source) for source in data_sources], ensure_ascii=False)


def _csv_rows_from_sources(data_sources: list[DataSource]) -> list[list[str]]:
    rows = [["label", "source", "data"]]
    for source in data_sources:
        data_value = source.formatted or source.value
        rows.append([source.label, source.source, data_value])
    return rows


def _deserialize_sources(raw: str) -> list[DataSource]:
    try:
        items = json.loads(raw)
    except json.JSONDecodeError:
        return []
    return [DataSource(**item) for item in items if isinstance(item, dict)]


def _normalize_sheet_title(title: str, fallback: str) -> str:
    cleaned = re.sub(r"[\\/*?:\\[\\]]", "_", title)[:31]
    return cleaned or fallback


def _write_sheet_data(sheet, source: DataSource) -> None:
    sheet.append(["label", source.label])
    sheet.append(["source", source.source])
    sheet.append([])
    raw = source.formatted or source.value
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        sheet.append(["data", raw])
        return

    if isinstance(parsed, list):
        if parsed and all(isinstance(item, dict) for item in parsed):
            headers: list[str] = []
            for item in parsed:
                for key in item.keys():
                    if key not in headers:
                        headers.append(key)
            sheet.append(headers)
            for item in parsed:
                row = [item.get(key, "") for key in headers]
                sheet.append(row)
            return
        if parsed and all(isinstance(item, list) for item in parsed):
            max_len = max((len(row) for row in parsed), default=0)
            for row in parsed:
                padded = list(row) + [""] * (max_len - len(row))
                sheet.append(padded)
            return
        sheet.append(["data", json.dumps(parsed, ensure_ascii=False)])
        return

    if isinstance(parsed, dict):
        sheet.append(["key", "value"])
        for key, value in parsed.items():
            sheet.append([key, json.dumps(value, ensure_ascii=False)])
        return

    sheet.append(["data", json.dumps(parsed, ensure_ascii=False)])


def _build_excel(data_sources: list[DataSource]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)

    for idx, source in enumerate(data_sources, start=1):
        title = _normalize_sheet_title(source.label, f"data_{idx}")
        sheet = workbook.create_sheet(title=title)
        _write_sheet_data(sheet, source)
        _style_sheet(sheet)
        _add_chart_to_sheet(sheet, source)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    header_font = Font(bold=True)

    if sheet.max_row >= 4:
        for cell in sheet[4]:
            cell.fill = header_fill
            cell.font = header_font

    for column_cells in sheet.columns:
        max_length = 0
        column = column_cells[0].column
        is_numeric = True
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            if not isinstance(value, (int, float)):
                is_numeric = False
            value_length = len(str(value))
            if value_length > max_length:
                max_length = value_length
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = min(max_length + 2, 50)
        if is_numeric:
            for cell in column_cells:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if isinstance(cell.value, float) else "#,##0"


def _detect_chart_type(source: DataSource, parsed: object) -> str:
    raw_text = (source.formatted or source.value).lower()
    if "type" in raw_text:
        if any(keyword in raw_text for keyword in ("pie", "piechart")):
            return "pie"
        if any(keyword in raw_text for keyword in ("bar", "column")):
            return "bar"
        if "line" in raw_text:
            return "line"

    if isinstance(parsed, list) and parsed:
        if all(isinstance(item, dict) for item in parsed):
            has_time = any(
                any("time" in key.lower() or "date" in key.lower() for key in item.keys())
                for item in parsed
            )
            if has_time:
                return "line"
            return "pie" if len(parsed) < 10 else "line"
        if all(isinstance(item, list) for item in parsed):
            return "pie" if len(parsed) < 10 else "line"
        return "pie" if len(parsed) < 10 else "line"
    if isinstance(parsed, dict):
        return "pie" if len(parsed) < 10 else "line"
    return "line"


def _add_chart_to_sheet(sheet, source: DataSource) -> None:
    raw = source.formatted or source.value
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        return

    data_start_row = 4
    max_row = sheet.max_row
    max_col = sheet.max_column
    if max_row < data_start_row + 1 or max_col < 1:
        return

    chart_type = _detect_chart_type(source, parsed)
    if chart_type == "pie":
        chart = PieChart()
    elif chart_type == "bar":
        chart = BarChart()
    else:
        chart = LineChart()

    if max_col >= 2:
        data_ref = Reference(sheet, min_col=2, min_row=data_start_row, max_row=max_row, max_col=max_col)
        categories = Reference(sheet, min_col=1, min_row=data_start_row + 1, max_row=max_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(categories)
    else:
        data_ref = Reference(sheet, min_col=1, min_row=data_start_row, max_row=max_row)
        chart.add_data(data_ref, titles_from_data=True)

    chart.height = 9
    chart.width = 16
    sheet.add_chart(chart, "F4")


def _write_docx_table(document: Document, source: DataSource) -> None:
    document.add_paragraph(f"数据源：{source.label}")
    document.add_paragraph(f"来源：{source.source}")
    raw = source.formatted or source.value
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError:
        document.add_paragraph(raw)
        document.add_page_break()
        return

    if isinstance(parsed, list) and parsed:
        if all(isinstance(item, dict) for item in parsed):
            headers: list[str] = []
            for item in parsed:
                for key in item.keys():
                    if key not in headers:
                        headers.append(key)
            table = document.add_table(rows=1, cols=len(headers))
            for idx, header in enumerate(headers):
                table.rows[0].cells[idx].text = str(header)
            for item in parsed:
                row_cells = table.add_row().cells
                for idx, header in enumerate(headers):
                    row_cells[idx].text = str(item.get(header, ""))
            document.add_page_break()
            return

        if all(isinstance(item, list) for item in parsed):
            max_len = max((len(row) for row in parsed), default=0)
            table = document.add_table(rows=0, cols=max_len or 1)
            for row in parsed:
                row_cells = table.add_row().cells
                padded = list(row) + [""] * (max_len - len(row))
                for idx, value in enumerate(padded):
                    row_cells[idx].text = str(value)
            document.add_page_break()
            return

    if isinstance(parsed, dict):
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "key"
        table.rows[0].cells[1].text = "value"
        for key, value in parsed.items():
            row_cells = table.add_row().cells
            row_cells[0].text = str(key)
            row_cells[1].text = json.dumps(value, ensure_ascii=False)
        document.add_page_break()
        return

    document.add_paragraph(json.dumps(parsed, ensure_ascii=False))
    document.add_page_break()


def _build_docx(data_sources: list[DataSource]) -> bytes:
    document = Document()
    for source in data_sources:
        _write_docx_table(document, source)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _prepare_form_data() -> dict[str, Any]:
    csrf_token = session.get("csrf_token")
    if not csrf_token:
        csrf_token = secrets.token_urlsafe(16)
        session["csrf_token"] = csrf_token
    return {
        "url": "",
        "cookies": "",
        "headers": "",
        "proxy": "",
        "timeout": str(DEFAULT_TIMEOUT),
        "custom_regex": "",
        "use_dynamic": "",
        "use_interactive": "",
        "data_sources": [],
        "raw_snippets": [],
        "data_sources_json": "",
        "error": "",
        "suggestions": [],
        "searched": False,
        "csrf_token": csrf_token,
    }


@app.route("/", methods=["GET", "POST"])
def index():
    data = _prepare_form_data()

    if request.method == "POST":
        data.update({key: request.form.get(key, "").strip() for key in data if isinstance(data[key], str)})
        data["searched"] = True

        if not data["csrf_token"] or data["csrf_token"] != session.get("csrf_token"):
            data["error"] = "安全校验失败，请刷新页面后重试。"
            logger.warning("CSRF token mismatch.")
            return render_template("index.html", **data)

        if not data["url"]:
            data["error"] = "请输入网址。"
            return render_template("index.html", **data)

        target_url = normalize_url(data["url"])
        parsed = urlparse(target_url)
        if not parsed.hostname or _is_private_host(parsed.hostname):
            data["error"] = "出于安全考虑，暂不支持访问内网或本地主机地址。"
            return render_template("index.html", **data)

        try:
            headers, cookies, proxies, timeout, custom_regex = _build_requests_options(data)
            if data.get("use_dynamic"):
                data_sources, raw_snippets = _scrape_dynamic_content(
                    target_url,
                    headers=headers,
                    cookies=cookies,
                    timeout=timeout,
                    custom_regex=custom_regex,
                    interactive=bool(data.get("use_interactive")),
                )
            else:
                data_sources, raw_snippets = scrape_content(
                    target_url,
                    headers=headers,
                    cookies=cookies,
                    proxies=proxies,
                    timeout=timeout,
                    custom_regex=custom_regex,
                )
            data["url"] = target_url
            data["data_sources"] = data_sources
            data["raw_snippets"] = raw_snippets
            data["data_sources_json"] = _serialize_data_sources(data_sources)
            if not data_sources:
                data["suggestions"] = [
                    "未找到图表数据，建议启用“动态爬取”后重试。",
                    "如果页面需要登录，请在高级设置中补充 Cookie 或 Headers。",
                    "若数据来自接口，可在浏览器 Network 中定位真实数据源。",
                ]
            logger.info("Scrape success: %s, data_sources=%s", target_url, len(data_sources))
        except ValueError as exc:
            data["error"] = str(exc)
            logger.warning("Input error: %s", exc)
        except requests.RequestException as exc:
            data["error"] = f"请求失败：{exc}"
            logger.warning("Request failed: %s", exc)
        except Exception as exc:  # noqa: BLE001
            data["error"] = f"解析失败：{exc}"
            logger.exception("Unexpected error")

    return render_template("index.html", **data)


@app.route("/export/json", methods=["POST"])
def export_json() -> Response:
    if request.form.get("csrf_token") != session.get("csrf_token"):
        return Response("Invalid CSRF token.", status=400)
    raw = request.form.get("data_sources_json", "")
    sources = _deserialize_sources(raw)
    payload = json.dumps([asdict(source) for source in sources], ensure_ascii=False, indent=2)
    return Response(
        payload,
        mimetype="application/json",
        headers={"Content-Disposition": "attachment; filename=chart-data.json"},
    )


@app.route("/export/csv", methods=["POST"])
def export_csv() -> Response:
    if request.form.get("csrf_token") != session.get("csrf_token"):
        return Response("Invalid CSRF token.", status=400)
    raw = request.form.get("data_sources_json", "")
    sources = _deserialize_sources(raw)
    output = io.StringIO()
    writer = csv.writer(output)
    for row in _csv_rows_from_sources(sources):
        writer.writerow(row)
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=chart-data.csv"},
    )


@app.route("/export/xlsx", methods=["POST"])
def export_xlsx() -> Response:
    if request.form.get("csrf_token") != session.get("csrf_token"):
        return Response("Invalid CSRF token.", status=400)
    raw = request.form.get("data_sources_json", "")
    sources = _deserialize_sources(raw)
    payload = _build_excel(sources)
    return Response(
        payload,
        mimetype=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": "attachment; filename=chart-data.xlsx"},
    )


@app.route("/export/docx", methods=["POST"])
def export_docx() -> Response:
    if request.form.get("csrf_token") != session.get("csrf_token"):
        return Response("Invalid CSRF token.", status=400)
    raw = request.form.get("data_sources_json", "")
    sources = _deserialize_sources(raw)
    payload = _build_docx(sources)
    return Response(
        payload,
        mimetype=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        headers={"Content-Disposition": "attachment; filename=chart-data.docx"},
    )


if __name__ == "__main__":
    _auto_install_requirements()
    def open_browser() -> None:
        webbrowser.open("http://localhost:5000")

    threading.Timer(1.0, open_browser).start()
    app.run(host="0.0.0.0", port=5000, debug=True)
